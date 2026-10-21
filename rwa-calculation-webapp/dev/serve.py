"""Local dev server for the RWA webapp front end — mocks the backend API with plausible
data and drives the Trace pages with the *real* ExcelDeepTraceEngine over a synthetic
RWA workbook.  Not deployed:  python dev/serve.py  ->  http://127.0.0.1:5178
"""

import io
import pathlib
import sys
import time

import openpyxl
from flask import Flask, request, jsonify, Response

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "python-lib"))
FRONTEND = ROOT / "frontend"

import excel_deep_trace as edt  # noqa: E402

app = Flask(__name__)


def ok(data=None, message=None, code=200):
    return jsonify({"success": True, "message": message, "data": data}), code


def bad(message, code=400, details=None):
    return jsonify({"success": False, "message": message, "details": details}), code


# --------------------------------------------------------------------------- fake state
CATS = {
    "corporate": {"display_name": "Corporate", "description": "Corporate exposure RWA",
                  "icon": "building", "recipe_id": "compute_QvDsUspv", "recipe_name": "Corporate RWA",
                  "folders": {"data": "mReNlOzQ", "mapping": "gCxSFlpK", "output": "5I5ZSrWc", "template": "iK1IJYr9"}},
    "institutions": {"display_name": "Institutions", "description": "Institutions exposure RWA",
                     "icon": "bank", "recipe_id": "compute_RDrI9CkK", "recipe_name": "",
                     "folders": {"data": "WB28JAfi", "mapping": "shveMrqs", "output": "cXkZJwZy", "template": "DAVd0l9d"}},
    "government": {"display_name": "Government", "description": "Government exposure RWA",
                   "icon": "landmark", "recipe_id": "", "recipe_name": "",
                   "folders": {"data": "", "mapping": "", "output": "", "template": ""}},
}
FILES = {
    ("corporate", "data"): [{"path": "/june_exposures.xlsx", "filename": "june_exposures.xlsx", "extension": ".xlsx", "size": 84210, "last_modified": "2026-09-01T09:12"},
                            {"path": "/july_exposures.xlsx", "filename": "july_exposures.xlsx", "extension": ".xlsx", "size": 91500, "last_modified": "2026-09-02T08:40"}],
    ("corporate", "mapping"): [{"path": "/rating_map.xlsx", "filename": "rating_map.xlsx", "extension": ".xlsx", "size": 12010, "last_modified": "2026-08-20T10:00"}],
    ("corporate", "output"): [{"path": "/Corporate_RWA_20260902_084500.xlsx", "filename": "Corporate_RWA_20260902_084500.xlsx", "extension": ".xlsx", "size": 120400, "last_modified": "2026-09-02T08:46"}],
    ("corporate", "template"): [{"path": "/Corporate_RWA_20260902_084500.xlsx", "filename": "Corporate_RWA_20260902_084500.xlsx", "extension": ".xlsx", "size": 120400, "last_modified": "2026-09-02T08:46"}],
    ("institutions", "data"): [],
    ("institutions", "mapping"): [],
    ("institutions", "output"): [],
    ("institutions", "template"): [],
}
ADMIN_TOKEN = "dev-token"


def _status(key):
    d = len(FILES.get((key, "data"), []))
    m = len(FILES.get((key, "mapping"), []))
    cfg = CATS[key]
    if not cfg["recipe_id"] or not cfg["folders"]["data"]:
        return "RECIPE_REQUIRED"
    if m != 1:
        return "MAPPING_REQUIRED"
    if d == 0:
        return "DATA_REQUIRED"
    return "READY"


@app.route("/api/whoami")
def whoami():
    return ok({"app_title": "RWA Calculation", "user": "yavnv"})


@app.route("/api/categories")
def categories():
    out = []
    for key, cfg in CATS.items():
        out.append({
            "key": key, "display_name": cfg["display_name"], "description": cfg["description"],
            "icon": cfg["icon"], "recipe_configured": bool(cfg["recipe_id"]),
            "data_file_count": len(FILES.get((key, "data"), [])),
            "mapping_file_count": len(FILES.get((key, "mapping"), [])),
            "output_file_count": len(FILES.get((key, "output"), [])),
            "template_file_count": len(FILES.get((key, "template"), [])),
            "status": _status(key),
            "split_enabled": False,
            "split_rules": [
                {"name": "COLLATERIZED", "column": "TRANCHE_COLLATERAL_TYPE", "value": "COLLATERIZED"},
                {"name": "NOT COLLATERIZED", "column": "TRANCHE_COLLATERAL_TYPE", "value": "NOT COLLATERIZED"},
                {"name": "GUARANTEED", "column": "TRANCHE_GUARANTEE_TYPE", "value": "GUARANTEED"},
            ],
        })
    return ok({"app_title": "RWA Calculation", "categories": out})


@app.route("/api/files")
def files():
    key, ft = request.args.get("category"), request.args.get("folder_type")
    return ok({"category": key, "folder_type": ft, "files": FILES.get((key, ft), []), "count": len(FILES.get((key, ft), []))})


@app.route("/api/upload", methods=["POST"])
def upload():
    key = request.form.get("category")
    ft = request.form.get("folder_type")
    got = request.files.getlist("files")
    recs = []
    if ft == "mapping":
        FILES[(key, "mapping")] = []
    for f in got:
        rec = {"path": "/" + f.filename, "filename": f.filename, "extension": "." + f.filename.rsplit(".", 1)[-1],
               "size": len(f.read()), "last_modified": time.strftime("%Y-%m-%dT%H:%M")}
        FILES.setdefault((key, ft), []).append(rec)
        recs.append(rec)
    return ok({"uploaded_files": recs}, "Uploaded.", 201)


@app.route("/api/delete", methods=["POST"])
def delete():
    b = request.get_json(force=True)
    lst = FILES.get((b["category"], b["folder_type"]), [])
    FILES[(b["category"], b["folder_type"])] = [x for x in lst if x["path"] != b["path"]]
    return ok(message="Deleted.")


@app.route("/api/download")
def download():
    return Response(b"mock file", headers={"Content-Disposition": 'attachment; filename="mock.bin"'})


SPLIT_COLUMNS = ["ID", "EXPOSURE", "TRANCHE_COLLATERAL_TYPE", "TRANCHE_GUARANTEE_TYPE", "RATING"]


@app.route("/api/split-preview", methods=["POST"])
def split_preview():
    b = request.get_json(force=True)
    paths = b.get("selected_paths") or [f["path"] for f in FILES.get((b["category"], "data"), [])]
    rules = b.get("rules") or []
    files = [{"filename": p.lstrip("/"), "rows": 1284,
              "datasets": [{"name": r["name"], "rows": None if r["column"] not in SPLIT_COLUMNS else 300 + 17 * i}
                           for i, r in enumerate(rules)]} for p in paths]
    return ok({"columns": SPLIT_COLUMNS, "files": files, "rules_error": None,
               "missing_columns": sorted({r["column"] for r in rules if r["column"] not in SPLIT_COLUMNS})})


@app.route("/api/validate-inputs", methods=["POST"])
def validate():
    b = request.get_json(force=True)
    n = len(b.get("selected_paths", []))
    return ok({
        "valid": n > 0, "category": b["category"],
        "input_row_count": 1284 * n, "input_column_count": 23,
        "mapping_row_count": 60, "mapping_column_count": 4,
        "mapping_file": {"filename": "rating_map.xlsx"},
        "input_files": [{"filename": p.lstrip("/"), "row_count": 1284, "column_count": 23, "detected_header_row": 2}
                        for p in b.get("selected_paths", [])],
        "errors": [] if n else ["No input data files were selected."],
    }, "Validation completed.")


@app.route("/api/run", methods=["POST"])
def run():
    b = request.get_json(force=True)
    return ok({
        "run_id": "RUN_20260902_" + time.strftime("%H%M%S"), "category": b["category"],
        "recipe_id": CATS[b["category"]]["recipe_id"], "status": "COMPLETED",
        "started_at": "2026-09-02T08:44:10", "completed_at": "2026-09-02T08:46:02",
        "output_file": "Corporate_RWA_20260902_084500.xlsx",
        "produced_files": [{"filename": "Corporate_RWA_20260902_084500.xlsx"}],
        "execution": {"state": "COMPLETED"},
        "split": None if not (b.get("split") or {}).get("enabled") else {
            "datasets": [{"source_filename": p.lstrip("/"), "split_name": r["name"], "column": r["column"],
                          "value": r["value"], "row_count": 321}
                         for p in b.get("selected_paths", []) for r in b["split"]["rules"]],
            "skipped": []},
    }, "Recipe completed.")


@app.route("/api/output-schema")
def output_schema():
    return ok({"category": request.args.get("category"),
               "columns": ["EXPOSURE", "RISK_WEIGHT", "TIER", "CRM", "CALCULATED_RWA", "SOURCE_RWA", "STATUS"],
               "file_count": 1, "detected_header_row": 1, "source_folder": "template"})


@app.route("/api/compare-all-outputs", methods=["POST"])
def compare():
    return ok({"category": request.get_json(force=True).get("category"), "source_folder": "template",
               "files_compared": 1, "total_rows": 1284, "matched_rows": 1201, "mismatched_rows": 83,
               "match_rate": 93.53, "warnings": [],
               "comparison_file": "Corporate_COMPARISON_20260902_090000.xlsx",
               "detected_headers": [{"file": "Corporate_RWA_20260902_084500.xlsx", "header_row": 1}],
               "mismatch_workbooks": ["Corporate_RWA_20260902_084500.xlsx"]},
              "Comparison completed.")


# ---- admin
@app.route("/api/login", methods=["POST"])
def login():
    b = request.get_json(force=True)
    if b.get("username") == "admin" and b.get("password") == "changeme":
        return ok({"token": ADMIN_TOKEN, "must_change": True})
    return bad("Invalid credentials.", 401)


@app.route("/api/admin/config")
def admin_config():
    if request.headers.get("X-Admin-Token") != ADMIN_TOKEN:
        return bad("Admin authentication required.", 401)
    return ok({"app_title": "RWA Calculation",
               "categories": [dict(key=k, **{kk: vv for kk, vv in v.items()}) for k, v in CATS.items()]})


@app.route("/api/admin/category/<key>", methods=["POST"])
def admin_save(key):
    if request.headers.get("X-Admin-Token") != ADMIN_TOKEN:
        return bad("Admin authentication required.", 401)
    b = request.get_json(force=True)
    CATS[key]["display_name"] = b.get("display_name") or CATS[key]["display_name"]
    CATS[key]["recipe_name"] = b.get("recipe_name", "")
    CATS[key]["recipe_id"] = b.get("recipe_id", "")
    CATS[key]["folders"] = b.get("folders", CATS[key]["folders"])
    return ok({"key": key}, "Saved.")


@app.route("/api/admin/change-credentials", methods=["POST"])
def admin_creds():
    return ok(message="Updated.")


# ---- trace (real engine over a synthetic workbook) --------------------------------
def _build_engine():
    ratings = [("AAA", 0.2), ("AA", 0.5), ("A", 1.0), ("BBB", 1.5), ("BB", 2.0)]
    rows = [(1_000_000, "AAA", 0.20, 200_000), (450_000, "A", 0.00, 450_000),
            (2_000_000, "BBB", 0.10, 2_650_000), (750_000, "AA", 0.25, 300_000),
            (1_200_000, "BB", 0.05, 2_280_000)]

    wf = openpyxl.Workbook()
    ws = wf.active
    ws.title = "RWA"
    ws.append(["EXPOSURE", "RISK_WEIGHT", "TIER", "CRM", "CALCULATED_RWA", "SOURCE_RWA", "STATUS"])
    for i, (exp, tier, crm, src) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=exp)
        ws.cell(row=i, column=2, value="=VLOOKUP(C{0},Ratings!$A$2:$B$6,2,FALSE)".format(i))
        ws.cell(row=i, column=3, value=tier)
        ws.cell(row=i, column=4, value=crm)
        ws.cell(row=i, column=5, value="=ROUND(A{0}*B{0}*(1-D{0}),2)".format(i))
        ws.cell(row=i, column=6, value=src)
        ws.cell(row=i, column=7, value='=IF(ABS(E{0}-F{0})<=0.01,"MATCH","MISMATCH")'.format(i))
    rt = wf.create_sheet("Ratings")
    rt.append(["TIER", "WEIGHT"])
    for t, w in ratings:
        rt.append([t, w])
    fb = io.BytesIO(); wf.save(fb)

    wv = openpyxl.Workbook()
    wsv = wv.active
    wsv.title = "RWA"
    wsv.append(["EXPOSURE", "RISK_WEIGHT", "TIER", "CRM", "CALCULATED_RWA", "SOURCE_RWA", "STATUS"])
    rd = dict(ratings)
    for i, (exp, tier, crm, src) in enumerate(rows, start=2):
        calc = round(exp * rd[tier] * (1 - crm), 2)
        vals = [exp, rd[tier], tier, crm, calc, src, "MATCH" if abs(calc - src) <= 0.01 else "MISMATCH"]
        for col, val in enumerate(vals, start=1):
            wsv.cell(row=i, column=col, value=val)
    rtv = wv.create_sheet("Ratings")
    rtv.append(["TIER", "WEIGHT"])
    for t, w in ratings:
        rtv.append([t, w])
    vb = io.BytesIO(); wv.save(vb)

    eng = edt.ExcelDeepTraceEngine(workbook_bytes=fb.getvalue(), filename="Corporate_RWA_20260902_084500.xlsx")
    eng._book.wb_values = openpyxl.load_workbook(io.BytesIO(vb.getvalue()), data_only=True)
    eng.wb_values = eng._book.wb_values
    eng.scan()
    return eng


SESS = {}


@app.route("/excel-trace/open-output", methods=["POST"])
def open_output():
    eng = _build_engine()
    sid = "sess-" + str(len(SESS) + 1)
    SESS[sid] = eng
    return jsonify({"success": True, "session_id": sid, "filename": eng.filename, "summary": eng.scan()})


def _eng():
    return SESS[request.get_json(force=True)["session_id"]]


@app.route("/excel-trace/row-journey", methods=["POST"])
def row_journey():
    b = request.get_json(force=True)
    return jsonify({"success": True, "journey": SESS[b["session_id"]].row_journey(b["sheet"], b["row"])})


@app.route("/excel-trace/expression-tree", methods=["POST"])
def expr_tree():
    b = request.get_json(force=True)
    return jsonify({"success": True, "tree": SESS[b["session_id"]].expression_tree(b["cell"])})


@app.route("/excel-trace/explain-cell", methods=["POST"])
def explain_cell():
    b = request.get_json(force=True)
    eng = SESS[b["session_id"]]
    key = eng._normalize_key(b["cell"])
    rec = eng.records.get(key)
    return jsonify({
        "success": True, "cell": key, "value": eng.value(key),
        "formula": rec.formula if rec else None,
        "formula_metadata": {
            "functions": rec.functions if rec else [], "formula_family": rec.family if rec else None,
            "dependency_count": rec.dependency_count if rec else 0,
        },
        "branch_analysis": eng.branch_analysis(key), "lookup_analysis": eng.lookup_analysis(key),
        "precedent_graph": eng.trace(key, "precedents", 8, 200),
        "dependent_graph": eng.trace(key, "dependents", 8, 200),
    })


@app.route("/excel-trace/findings")
def findings():
    eng = SESS[request.args.get("session_id")]
    return jsonify({"success": True, "total": len(eng.findings),
                    "findings": [{"severity": f.severity, "code": f.code, "cell": f.cell,
                                  "message": f.message, "evidence": f.evidence} for f in eng.findings]})


@app.route("/excel-trace/columns", methods=["POST"])
def columns():
    b = request.get_json(force=True)
    return jsonify({"success": True, **SESS[b["session_id"]].sheet_columns(b["sheet"])})


@app.route("/excel-trace/mismatch-rows", methods=["POST"])
def mismatch_rows():
    b = request.get_json(force=True)
    return jsonify({"success": True, **SESS[b["session_id"]].mismatch_rows(
        b["sheet"], b["left"], b["right"], float(b.get("tolerance", 0.01)))})


@app.route("/")
def index():
    body = (FRONTEND / "body.html").read_text(encoding="utf-8")
    css = (FRONTEND / "style.css").read_text(encoding="utf-8")
    js = (FRONTEND / "app.js").read_text(encoding="utf-8")
    return Response(
        f"""<!doctype html><html><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>RWA Calculation (dev)</title><style>{css}</style></head>
<body>{body}<script>{js}</script></body></html>""",
        mimetype="text/html")


if __name__ == "__main__":
    app.run(port=5178, debug=False)
