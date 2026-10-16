"""Fixtures: a small RWA-style workbook with live formulas + a matching values
workbook (openpyxl can't calculate, so we inject the cached values ourselves)."""

import io
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "python-lib"))


def _save(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


# layout, row by row: [Exposure, RiskWeight, Tier, CRM, CALCULATED_RWA, SOURCE_RWA, STATUS]
_ROWS = [
    (1000, "AAA", 0.2, 200),    # RW .2  -> 1000*.2*(1-.2) = 160   vs 200 -> MISMATCH
    (500, "A", 0.0, 500),       # RW 1.0 -> 500*1*(1-0)   = 500   vs 500 -> MATCH
    (2000, "BBB", 0.1, 2700),   # RW 1.5 -> 2000*1.5*(1-.1) = 2700 vs 2700 -> MATCH
]
_RATINGS = [("AAA", 0.2), ("AA", 0.5), ("A", 1.0), ("BBB", 1.5)]


def _rw(tier):
    return dict(_RATINGS)[tier]


@pytest.fixture
def rwa_engine():
    from excel_deep_trace import ExcelDeepTraceEngine

    # --- formula workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RWA"
    ws.append(["Exposure", "RiskWeight", "Tier", "CRM", "CALCULATED_RWA", "SOURCE_RWA", "STATUS"])
    for i, (exp, tier, crm, src) in enumerate(_ROWS, start=2):
        ws.cell(row=i, column=1, value=exp)
        ws.cell(row=i, column=2, value="=VLOOKUP(C{0},Ratings!$A$2:$B$5,2,FALSE)".format(i))
        ws.cell(row=i, column=3, value=tier)
        ws.cell(row=i, column=4, value=crm)
        ws.cell(row=i, column=5, value="=A{0}*B{0}*(1-D{0})".format(i))
        ws.cell(row=i, column=6, value=src)
        ws.cell(row=i, column=7,
                value='=IF(ABS(E{0}-F{0})<=0.01,"MATCH","MISMATCH")'.format(i))
    rt = wb.create_sheet("Ratings")
    rt.append(["Tier", "Weight"])
    for t, w in _RATINGS:
        rt.append([t, w])
    formula_bytes = _save(wb)

    # --- values workbook (same layout, formulas replaced with computed results)
    wv = openpyxl.Workbook()
    wsv = wv.active
    wsv.title = "RWA"
    wsv.append(["Exposure", "RiskWeight", "Tier", "CRM", "CALCULATED_RWA", "SOURCE_RWA", "STATUS"])
    for i, (exp, tier, crm, src) in enumerate(_ROWS, start=2):
        rw = _rw(tier)
        calc = exp * rw * (1 - crm)
        wsv.cell(row=i, column=1, value=exp)
        wsv.cell(row=i, column=2, value=rw)
        wsv.cell(row=i, column=3, value=tier)
        wsv.cell(row=i, column=4, value=crm)
        wsv.cell(row=i, column=5, value=calc)
        wsv.cell(row=i, column=6, value=src)
        wsv.cell(row=i, column=7, value="MATCH" if abs(calc - src) <= 0.01 else "MISMATCH")
    rtv = wv.create_sheet("Ratings")
    rtv.append(["Tier", "Weight"])
    for t, w in _RATINGS:
        rtv.append([t, w])
    values_bytes = _save(wv)

    eng = ExcelDeepTraceEngine(workbook_bytes=formula_bytes, filename="june_rwa.xlsx")
    eng._book.wb_values = openpyxl.load_workbook(io.BytesIO(values_bytes), data_only=True)
    eng.wb_values = eng._book.wb_values
    eng.scan()
    return eng
